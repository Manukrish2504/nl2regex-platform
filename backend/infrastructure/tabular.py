"""Read CSV and Excel files as rows of text. Every cell stays text: nothing is type-converted."""

import csv
import io
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from domain.errors import ColumnNotFound, FileUnreadable
from domain.tables import Row

# An .xlsx is a zip archive that must be downloaded whole before a single row can be read.
EXCEL_MAX_BYTES = 200 * 1024 * 1024
# Enough of a CSV to learn its columns, delimiter and first rows without downloading the rest.
CSV_HEAD_BYTES = 256 * 1024
_SNIFF_BYTES = 64 * 1024
_DELIMITERS = ",;\t|"

# The default cap of 128 KB per cell rejects legitimate long free-text fields.
csv.field_size_limit(16 * 1024 * 1024)


def sniff_delimiter(head: bytes) -> str:
    """The delimiter of a CSV, judged from its first bytes."""
    sample = head[:_SNIFF_BYTES].decode("utf-8-sig", errors="ignore")
    try:
        # Only the delimiter is taken from the sniffer. It also guesses the quoting rules, and
        # gets them wrong often enough to mangle quoted fields, so those stay standard.
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        return ","  # one column, or too little to go on: a comma is the safe guess


def csv_header(head: bytes) -> tuple[list[str], str]:
    """Column names and delimiter from the first bytes of a CSV, for engines that read the rest
    themselves. `head` must end on a line boundary."""
    try:
        head.decode("utf-8-sig")  # strict, so a wrong encoding fails here and not as mojibake
    except UnicodeDecodeError as error:
        raise FileUnreadable(
            "The file is not UTF-8 encoded. Save it as UTF-8 and try again."
        ) from error
    return read_header(csv_rows(io.BufferedReader(io.BytesIO(head)))), sniff_delimiter(head)


def csv_rows(stream: io.BufferedIOBase) -> Iterator[Row]:
    """Rows of a UTF-8 CSV (with or without a byte-order mark), with its delimiter detected."""
    delimiter = sniff_delimiter(stream.peek(_SNIFF_BYTES))
    text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    try:
        for row in csv.reader(text, delimiter=delimiter):
            if not row:
                continue  # a blank line is not a row; Spark's CSV reader skips them too
            # CSV has no way to say "no value", so an empty cell is read as None. Spark's CSV
            # reader does the same by default, which keeps every engine's output identical.
            yield [cell or None for cell in row]
    except UnicodeDecodeError as error:
        raise FileUnreadable(
            "The file is not UTF-8 encoded. Save it as UTF-8 and try again."
        ) from error
    except csv.Error as error:
        raise FileUnreadable(f"The CSV is malformed: {error}.") from error


def excel_rows(path: Path) -> Iterator[Row]:
    """Rows of the first sheet. Cells become text; rows that are completely empty are skipped."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for a corrupt file
        raise FileUnreadable("The Excel file could not be opened. It may be corrupt.") from error
    try:
        for cells in workbook.worksheets[0].iter_rows(values_only=True):
            if any(cell is not None for cell in cells):
                yield [None if cell is None else str(cell) for cell in cells]
    finally:
        workbook.close()


def read_header(rows: Iterator[Row]) -> list[str]:
    """Consume the first row as column names, made unique and non-empty."""
    header = next(rows, None)
    if header is None:
        raise FileUnreadable("The file is empty.")
    names, seen = [], {}
    for position, cell in enumerate(header, start=1):
        name = (cell or "").strip() or f"column_{position}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    if not names:
        raise FileUnreadable("The file has no header row.")
    return names


def fit(row: Row, width: int) -> Row:
    """Make a row match the header: short rows are padded with None, extra cells are dropped.

    This is what Spark's CSV reader does in its default (permissive) mode, so every engine agrees.
    """
    return row[:width] + [None] * (width - len(row))


def target_indexes(header: list[str], wanted: list[str]) -> list[int]:
    """Positions of the wanted columns. An empty request means every column.

    An exact name wins. Failing that, case, spaces, underscores and hyphens are ignored, so a
    request that says "the email column" finds `Email`, and "first name" finds `First_Name`. A
    name that still matches nothing fails the job; it never quietly means "every column".
    """
    if not wanted:
        return list(range(len(header)))
    loosely = {_loose(name): index for index, name in reversed(list(enumerate(header)))}
    indexes, missing = [], []
    for name in wanted:
        index = header.index(name) if name in header else loosely.get(_loose(name))
        if index is None:
            missing.append(name)
        elif index not in indexes:
            indexes.append(index)
    if missing:
        raise ColumnNotFound(missing)
    return indexes


def _loose(name: str) -> str:
    return "".join(character for character in name.casefold() if character not in " _-")
